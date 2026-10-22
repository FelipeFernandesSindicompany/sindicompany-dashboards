"""
Lista os achados que o responsável mandou RETIRAR DO RELATÓRIO (PDF) mas que continuam registrados pelo motor.

"Retirar" = não aparece no PDF entregue à administradora. A validação do motor segue inteira: as regras rodam, o achado fica em
data/<condomínio>/conciliacao/<mês>/<versão>/achados_brutos.json e achados_revisados.json (revisado_por = "humano:dispensa", com a decisão e o
motivo em motivo_divergencia_da_sugestao). A decisão em si fica versionada em config/dispensas_validacao/<condomínio>/<AAAA-MM>.json.

Uso: python scripts/listar_achados_retirados.py      -> output/achados_retirados_do_relatorio_<data>.xlsx
"""
import glob
import json
import os
import sys
from datetime import datetime
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

sys.stdout.reconfigure(encoding="utf-8")
RAIZ = Path(__file__).resolve().parent.parent


def main() -> None:
    nomes = {c["id"]: c["nome"] for c in json.load(open(RAIZ / "config" / "condominios.json", encoding="utf-8"))["condominios"]}
    linhas = []
    for d in sorted(glob.glob(str(RAIZ / "data" / "*" / "conciliacao" / "20*"))):
        lat = os.path.join(d, "latest")
        if os.path.isdir(lat):
            versao = lat
        elif os.path.isfile(lat):
            versao = os.path.join(d, open(lat).read().strip())
        else:
            continue
        f_rev, f_bru = os.path.join(versao, "achados_revisados.json"), os.path.join(versao, "achados_brutos.json")
        if not (os.path.exists(f_rev) and os.path.exists(f_bru)):
            continue
        rev = json.load(open(f_rev, encoding="utf-8"))
        brutos = {x["id"]: x for x in json.load(open(f_bru, encoding="utf-8"))}
        partes = Path(d).parts
        cid, mes = partes[-3], partes[-1]
        for r in rev:
            if r.get("revisado_por") != "humano:dispensa":
                continue
            x = brutos.get(r["achado_id"], {})
            linhas.append([nomes.get(cid, cid), f"{mes[5:]}/{mes[:4]}", r["achado_id"], x.get("tipo"), x.get("regra_aplicada"), x.get("valor_encontrado"),
                           r.get("titulo"), r.get("motivo_divergencia_da_sugestao"), (r.get("revisado_em") or "")[:10]])
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Retirados do relatório"
    ws.append(["Condomínio", "Mês", "Achado", "Tipo", "Regra do motor", "Valor", "Título (como estava no relatório)", "Decisão / motivo", "Data"])
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="305496")
    for l in linhas:
        ws.append(l)
    for larg, col in zip([28, 9, 10, 26, 44, 12, 70, 80, 11], "ABCDEFGHI"):
        ws.column_dimensions[col].width = larg
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
        row[5].number_format = '#,##0.00;[Red]-#,##0.00'
    ws.freeze_panes = "A2"
    nota = wb.create_sheet("Como ler")
    nota.append(["Retirar um item = ele não aparece no PDF entregue. O motor continua validando: o achado segue gravado em achados_brutos.json e achados_revisados.json "
                 "da versão (revisado_por = humano:dispensa) e a decisão fica em config/dispensas_validacao/<condomínio>/<mês>.json (versionada no Git)."])
    nota.column_dimensions["A"].width = 140
    nota["A1"].alignment = Alignment(wrap_text=True, vertical="top")
    saida = RAIZ / "output" / f"achados_retirados_do_relatorio_{datetime.now():%Y-%m-%d}.xlsx"
    wb.save(saida)
    print(len(linhas), "achado(s) retirado(s) do relatório ->", saida)


if __name__ == "__main__":
    main()
